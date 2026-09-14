#!/usr/bin/env python3
"""
Generates the complete, perfectly-aligned 8-slide presentation with diagrams.
Adheres strictly to minimal design:
- Light clean backgrounds
- Charcoal/dark grey typography
- Restrained navy accent
- Perfectly aligned diagrams for minimal, intuitive understanding
- Zero overlapping, high whitespace
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

OUTPUT_PPTX = "docs/presentation/GLIDE_SPEC_40_CoFounder_Intro.pptx"
IMG_DIR = "docs/presentation/images"

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme colors
    COLOR_BG = RGBColor(255, 255, 255)
    COLOR_CARD = RGBColor(250, 250, 250)
    COLOR_BORDER = RGBColor(229, 231, 235)
    COLOR_MAIN = RGBColor(17, 24, 39)       # #111827
    COLOR_BODY = RGBColor(55, 65, 81)       # #374151
    COLOR_MUTED = RGBColor(107, 114, 128)   # #6B7280
    COLOR_LIGHT = RGBColor(156, 163, 175)   # #9CA3AF
    COLOR_ACCENT = RGBColor(30, 58, 138)    # #1E3A8A
    COLOR_ACCENT_BG = RGBColor(241, 245, 249)

    FONT_FAMILY = "Apple SD Gothic Neo"

    def apply_style(p, size=13, bold=False, color=COLOR_BODY, align=PP_ALIGN.LEFT, font_name=FONT_FAMILY):
        p.font.name = font_name
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
        p.alignment = align

    def add_slide_header(slide, meta_text, title_text, subtitle_text, slide_index_str):
        # Meta category
        box_meta = slide.shapes.add_textbox(Inches(0.8), Inches(0.48), Inches(11.7), Inches(0.28))
        tf_m = box_meta.text_frame
        tf_m.margin_left = tf_m.margin_top = tf_m.margin_right = tf_m.margin_bottom = 0
        apply_style(tf_m.paragraphs[0], size=10.5, bold=True, color=COLOR_LIGHT)
        tf_m.paragraphs[0].text = meta_text

        # Title
        box_t = slide.shapes.add_textbox(Inches(0.8), Inches(0.82), Inches(11.7), Inches(0.55))
        tf_t = box_t.text_frame
        tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
        apply_style(tf_t.paragraphs[0], size=24, bold=True, color=COLOR_MAIN)
        tf_t.paragraphs[0].text = title_text

        # Subtitle
        box_sub = slide.shapes.add_textbox(Inches(0.8), Inches(1.42), Inches(11.7), Inches(0.38))
        tf_sub = box_sub.text_frame
        tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
        apply_style(tf_sub.paragraphs[0], size=13, bold=False, color=COLOR_MUTED)
        tf_sub.paragraphs[0].text = subtitle_text

        # Divider line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.9), Inches(11.733), Inches(0.012))
        line.fill.solid()
        line.fill.fore_color.rgb = COLOR_BORDER
        line.line.color.rgb = COLOR_BORDER

        # Footer
        box_f = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.25))
        tf_f = box_f.text_frame
        tf_f.margin_left = tf_f.margin_top = tf_f.margin_right = tf_f.margin_bottom = 0
        apply_style(tf_f.paragraphs[0], size=9.5, bold=False, color=COLOR_LIGHT)
        tf_f.paragraphs[0].text = f"GLIDE-SPEC 40  ·  Co-founder discussion  ·  {slide_index_str}"

    # -------------------------------------------------------------
    # SLIDE 1: Title
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    
    b_tag = s1.shapes.add_textbox(Inches(1.2), Inches(1.9), Inches(10), Inches(0.35))
    tf_tag = b_tag.text_frame
    tf_tag.margin_left = tf_tag.margin_top = 0
    apply_style(tf_tag.paragraphs[0], size=11, bold=True, color=COLOR_MUTED)
    tf_tag.paragraphs[0].text = "PROJECT INTRODUCTION"

    b_title = s1.shapes.add_textbox(Inches(1.2), Inches(2.3), Inches(10), Inches(1.1))
    tf_title = b_title.text_frame
    tf_title.margin_left = tf_title.margin_top = 0
    apply_style(tf_title.paragraphs[0], size=48, bold=True, color=COLOR_MAIN)
    tf_title.paragraphs[0].text = "GLIDE-SPEC 40"

    b_sub = s1.shapes.add_textbox(Inches(1.2), Inches(3.55), Inches(10), Inches(0.45))
    tf_sub = b_sub.text_frame
    tf_sub.margin_left = tf_sub.margin_top = 0
    apply_style(tf_sub.paragraphs[0], size=20, bold=False, color=COLOR_BODY)
    tf_sub.paragraphs[0].text = "Technical Anti-Chafing Stick"

    div1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(4.25), Inches(4.5), Inches(0.015))
    div1.fill.solid()
    div1.fill.fore_color.rgb = COLOR_MAIN
    div1.line.color.rgb = COLOR_MAIN

    b_spec = s1.shapes.add_textbox(Inches(1.2), Inches(4.5), Inches(10), Inches(1.0))
    tf_spec = b_spec.text_frame
    tf_spec.margin_left = tf_spec.margin_top = 0
    p1 = tf_spec.paragraphs[0]
    apply_style(p1, size=15, bold=True, color=COLOR_MAIN)
    p1.text = "20g Powder-in-Balm Stick"
    p1.space_after = Pt(6)

    p2 = tf_spec.add_paragraph()
    apply_style(p2, size=13.5, bold=False, color=COLOR_MUTED)
    p2.text = "Marathon  ·  Trail  ·  Long-distance  ·  Military Ruck March"

    b_foot = s1.shapes.add_textbox(Inches(1.2), Inches(6.5), Inches(10), Inches(0.4))
    tf_foot = b_foot.text_frame
    tf_foot.margin_left = tf_foot.margin_top = 0
    p_f = tf_foot.paragraphs[0]
    apply_style(p_f, size=10, bold=False, color=COLOR_LIGHT)
    p_f.text = "Development baseline: Rev.7.3  |  Engineering status: Phase C — Physical Pilot Execution\nCo-founder discussion baseline"

    # -------------------------------------------------------------
    # SLIDE 2: 02 / PROBLEM (Text Left + Diagram Right)
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    add_slide_header(s2, "02 / PROBLEM", "장시간 활동에서 반복되는 피부 마찰 문제",
                     "제품 개발에서 해결해야 할 것은 단순한 윤활감이 아니라 실제 사용 중의 균형입니다.", "02 / 08")

    # Left Column (Text cards)
    col_w = Inches(4.8)
    card1 = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(2.15), col_w, Inches(1.4))
    card1.fill.solid()
    card1.fill.fore_color.rgb = COLOR_CARD
    card1.line.color.rgb = COLOR_BORDER
    tb1 = s2.shapes.add_textbox(Inches(0.95), Inches(2.25), col_w - Inches(0.3), Inches(1.2))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    apply_style(p, size=13, bold=True, color=COLOR_MAIN)
    p.text = "주요 사용 부위"
    p.space_after = Pt(4)
    p_sub = tf1.add_paragraph()
    apply_style(p_sub, size=11.5, color=COLOR_BODY)
    p_sub.text = "•  허벅지 안쪽  ·  사타구니  ·  겨드랑이  ·  발 & 뒤꿈치"

    card2 = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(3.7), col_w, Inches(1.5))
    card2.fill.solid()
    card2.fill.fore_color.rgb = COLOR_CARD
    card2.line.color.rgb = COLOR_BORDER
    tb2 = s2.shapes.add_textbox(Inches(0.95), Inches(3.8), col_w - Inches(0.3), Inches(1.3))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    apply_style(p, size=13, bold=True, color=COLOR_MAIN)
    p.text = "개발에서 보는 사용성 요건"
    p.space_after = Pt(4)
    p_sub = tf2.add_paragraph()
    apply_style(p_sub, size=11.5, color=COLOR_BODY)
    p_sub.text = "•  마찰 감소  ·  지속성  ·  사용감(보송함)  ·  의류 전이 최소화"

    card3 = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(5.35), col_w, Inches(1.4))
    card3.fill.solid()
    card3.fill.fore_color.rgb = COLOR_CARD
    card3.line.color.rgb = COLOR_BORDER
    tb3 = s2.shapes.add_textbox(Inches(0.95), Inches(5.45), col_w - Inches(0.3), Inches(1.2))
    tf3 = tb3.text_frame
    tf3.word_wrap = True
    p = tf3.paragraphs[0]
    apply_style(p, size=13, bold=True, color=COLOR_MAIN)
    p.text = "타깃 활동 환경"
    p.space_after = Pt(4)
    p_sub = tf3.add_paragraph()
    apply_style(p_sub, size=11.5, color=COLOR_BODY)
    p_sub.text = "•  러닝 풀코스  ·  트레일러닝  ·  장거리 보행  ·  20kg 완전군장 행군"

    # Right Column: Diagram Image
    img2_path = os.path.join(IMG_DIR, "diagram_friction_interface.png")
    if os.path.exists(img2_path):
        s2.shapes.add_picture(img2_path, Inches(5.9), Inches(2.15), Inches(6.6), Inches(3.6))

    # Callout below diagram
    callout2 = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.9), Inches(5.9), Inches(6.6), Inches(0.85))
    callout2.fill.solid()
    callout2.fill.fore_color.rgb = COLOR_ACCENT_BG
    callout2.line.color.rgb = COLOR_BORDER
    tb_c2 = s2.shapes.add_textbox(Inches(6.1), Inches(6.0), Inches(6.2), Inches(0.65))
    tf_c2 = tb_c2.text_frame
    tf_c2.word_wrap = True
    p_c2 = tf_c2.paragraphs[0]
    apply_style(p_c2, size=12, bold=True, color=COLOR_MAIN)
    p_c2.text = "핵심: 반복 마찰을 줄이면서도 땀·움직임·의류 접촉을 견디는 것."

    # -------------------------------------------------------------
    # SLIDE 3: 03 / PRODUCT (4 Cards Top + Mechanism Diagram Bottom)
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    add_slide_header(s3, "03 / PRODUCT", "제품은 4개의 기능 시스템으로 설계",
                     "Rev.7.3은 최종 생산 Formula가 아니라 Pilot으로 연결하기 위한 개발 기준선입니다.", "03 / 08")

    # 4 Cards
    w_card = Inches(2.78)
    gap_card = Inches(0.2)
    start_x = Inches(0.8)
    y_card = Inches(2.15)
    h_card = Inches(1.65)

    systems_data = [
        ("Powder", "마찰 저감 · 표면감", "구상 파우더가 표면 마찰을 낮추고 땀·유분을 흡착"),
        ("Wax", "스틱 구조 · 사용성", "상온 형상을 지지하고 도포량과 경도를 제어"),
        ("Silicone", "Spread · Slip", "얇고 매끄럽게 펴지도록 초기 슬립과 감촉을 형성"),
        ("Resin", "지속성 · 필름", "땀·수분에 견디는 표면막을 형성해 윤활층 유지")
    ]

    for idx, (s_name, s_role, s_desc) in enumerate(systems_data):
        x = start_x + idx * (w_card + gap_card)
        bg = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y_card, w_card, h_card)
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_CARD
        bg.line.color.rgb = COLOR_BORDER

        tb = s3.shapes.add_textbox(x + Inches(0.18), y_card + Inches(0.15), w_card - Inches(0.36), h_card - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = 0
        p = tf.paragraphs[0]
        apply_style(p, size=15, bold=True, color=COLOR_MAIN)
        p.text = s_name
        p.space_after = Pt(2)

        p_r = tf.add_paragraph()
        apply_style(p_r, size=11, bold=True, color=COLOR_ACCENT)
        p_r.text = f"→ {s_role}"
        p_r.space_after = Pt(6)

        p_d = tf.add_paragraph()
        apply_style(p_d, size=10.5, color=COLOR_BODY)
        p_d.text = s_desc

    # Diagram Image Bottom
    img3_path = os.path.join(IMG_DIR, "diagram_4system_mechanism.png")
    if os.path.exists(img3_path):
        s3.shapes.add_picture(img3_path, Inches(1.8), Inches(4.0), Inches(9.7), Inches(2.65))

    # Note
    tb_n3 = s3.shapes.add_textbox(Inches(0.8), Inches(6.7), Inches(11.7), Inches(0.3))
    tf_n3 = tb_n3.text_frame
    tf_n3.margin_left = tf_n3.margin_top = 0
    p_n3 = tf_n3.paragraphs[0]
    apply_style(p_n3, size=10, color=COLOR_MUTED)
    p_n3.text = "* 세부 원료 목록과 개별 배합비는 이 자료에서 생략하고, 시스템별 기능과 물성 제어에 집중합니다."

    # -------------------------------------------------------------
    # SLIDE 4: 04 / DEVELOPMENT APPROACH (Diagram Top + Principles Bottom)
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    add_slide_header(s4, "04 / DEVELOPMENT APPROACH", "제품보다 개발 과정에 더 큰 차이가 있다",
                     "원료 → 제조 → QC → 데이터가 이어지는 구조를 먼저 만든다.", "04 / 08")

    img4_path = os.path.join(IMG_DIR, "diagram_engineering_loop.png")
    if os.path.exists(img4_path):
        s4.shapes.add_picture(img4_path, Inches(0.8), Inches(2.15), Inches(11.733), Inches(2.5))

    # Principles card bottom
    card_p = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(4.85), Inches(11.733), Inches(1.9))
    card_p.fill.solid()
    card_p.fill.fore_color.rgb = COLOR_CARD
    card_p.line.color.rgb = COLOR_BORDER

    tb_p = s4.shapes.add_textbox(Inches(1.1), Inches(5.0), Inches(11.1), Inches(1.6))
    tf_p = tb_p.text_frame
    tf_p.word_wrap = True
    p_pt = tf_p.paragraphs[0]
    apply_style(p_pt, size=14, bold=True, color=COLOR_ACCENT)
    p_pt.text = "개발 원칙"
    p_pt.space_after = Pt(8)

    principles = [
        "원료 Lot 실측치 / 제조 공정 변수(온도·속도) / QC 측정값을 표준 스키마로 함께 기록",
        "정량 QC(경도, 전이도)와 통계적 실험계획법(DOE)을 통해 배치 간 차이와 오차를 설명",
        "실측 데이터가 쌓이면 다음 Formula를 시행착오 없이 체계적으로 최적화하는 폐루프 구축"
    ]
    for item in principles:
        pi = tf_p.add_paragraph()
        apply_style(pi, size=12, color=COLOR_BODY)
        pi.text = f"•  {item}"
        pi.space_after = Pt(4)

    # -------------------------------------------------------------
    # SLIDE 5: 05 / CURRENT STATUS (Status Cards Top + Stage Gate Bottom)
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    add_slide_header(s5, "05 / CURRENT STATUS", "현재는 실제 Pilot을 준비하는 단계",
                     "엔지니어링 프레임워크는 구축됐고, 물리적 제조와 통계적 Qualification이 남아 있습니다.", "05 / 08")

    # Top Left: 완료
    card_c = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(2.15), Inches(5.7), Inches(2.35))
    card_c.fill.solid()
    card_c.fill.fore_color.rgb = COLOR_CARD
    card_c.line.color.rgb = COLOR_BORDER

    tb_c = s5.shapes.add_textbox(Inches(1.0), Inches(2.25), Inches(5.3), Inches(2.1))
    tf_c = tb_c.text_frame
    tf_c.word_wrap = True
    p_ct = tf_c.paragraphs[0]
    apply_style(p_ct, size=13.5, bold=True, color=COLOR_MAIN)
    p_ct.text = "완료된 엔지니어링 자산"
    p_ct.space_after = Pt(6)

    dones = [
        "Rev.7.3 development baseline 확정",
        "원료 / 배합 데이터 구조 및 QC 측정 스키마",
        "DOE + Regression / Qualification framework",
        "Optimization framework & Revision tracking",
        "기초 시스템 테스트: 31 / 31 unit tests pass"
    ]
    for d in dones:
        pi = tf_c.add_paragraph()
        apply_style(pi, size=11, color=COLOR_BODY)
        pi.text = f"•  {d}"
        pi.space_after = Pt(2)

    # Top Right: 현재 및 다음
    card_n = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(2.15), Inches(5.7), Inches(2.35))
    card_n.fill.solid()
    card_n.fill.fore_color.rgb = COLOR_CARD
    card_n.line.color.rgb = COLOR_BORDER

    tb_n = s5.shapes.add_textbox(Inches(7.0), Inches(2.25), Inches(5.3), Inches(2.1))
    tf_n = tb_n.text_frame
    tf_n.word_wrap = True
    p_nt = tf_n.paragraphs[0]
    apply_style(p_nt, size=13.5, bold=True, color=COLOR_ACCENT)
    p_nt.text = "현재 진행 위치"
    p_nt.space_after = Pt(4)

    p_curr = tf_n.add_paragraph()
    apply_style(p_curr, size=12.5, bold=True, color=COLOR_MAIN)
    p_curr.text = "Phase C — Physical Pilot Execution"
    p_curr.space_after = Pt(2)

    p_cdesc = tf_n.add_paragraph()
    apply_style(p_cdesc, size=11, color=COLOR_MUTED)
    p_cdesc.text = "실제 원료를 계량·제조하고 QC 실측값을 확보하는 단계"
    p_cdesc.space_after = Pt(10)

    p_next = tf_n.add_paragraph()
    apply_style(p_next, size=12, bold=True, color=COLOR_MAIN)
    p_next.text = "다음 직전 단계:"
    p_next.space_after = Pt(2)

    p_ndesc = tf_n.add_paragraph()
    apply_style(p_ndesc, size=11, color=COLOR_BODY)
    p_ndesc.text = "실제 원료 입고  →  18-run Pilot  →  QC 실측  →  통계적 Qualification"

    # Bottom Diagram: Stage Gate Pipeline
    img5_path = os.path.join(IMG_DIR, "diagram_project_stage_gate.png")
    if os.path.exists(img5_path):
        s5.shapes.add_picture(img5_path, Inches(0.8), Inches(4.7), Inches(11.733), Inches(2.2))

    # -------------------------------------------------------------
    # SLIDE 6: 06 / NEXT STEPS (Roadmap)
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    add_slide_header(s6, "06 / NEXT STEPS", "이제부터는 실제 실행이 중요하다",
                     "Pilot 데이터가 확보되어야 제품과 Production Model을 함께 판단할 수 있습니다.", "06 / 08")

    steps_7 = [
        ("01", "원료 공급사", "CoA / 샘플 / Grade 확보"),
        ("02", "제조·Pilot 파트너", "소량 제조·시험 생산 설비 확보"),
        ("03", "18-run Pilot", "설계된 18개 배치 정밀 실행"),
        ("04", "QC + 사용자 테스트", "물성 실측 + 실제 필드 사용"),
        ("05", "분석·최적화", "회귀 / Formula refinement"),
        ("06", "Confirmation Run", "최종 재현성 검증"),
        ("07", "초기 생산 검토", "3,000개 목표 (COGS ~₩2,950)")
    ]

    w_s6 = Inches(1.53)
    gap_s6 = Inches(0.17)
    start_s6 = Inches(0.8)
    y_s6 = Inches(2.4)
    h_s6 = Inches(3.6)

    for i, (num, stitle, sdesc) in enumerate(steps_7):
        x = start_s6 + i * (w_s6 + gap_s6)
        card = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y_s6, w_s6, h_s6)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_BORDER

        tb = s6.shapes.add_textbox(x + Inches(0.12), y_s6 + Inches(0.18), w_s6 - Inches(0.24), h_s6 - Inches(0.36))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = 0

        p = tf.paragraphs[0]
        apply_style(p, size=13, bold=True, color=COLOR_ACCENT)
        p.text = num
        p.space_after = Pt(8)

        p_t = tf.add_paragraph()
        apply_style(p_t, size=13, bold=True, color=COLOR_MAIN)
        p_t.text = stitle
        p_t.space_after = Pt(10)

        p_d = tf.add_paragraph()
        apply_style(p_d, size=11, color=COLOR_MUTED)
        p_d.text = sdesc

    # Bottom notice
    tb_n6 = s6.shapes.add_textbox(Inches(0.8), Inches(6.35), Inches(11.7), Inches(0.35))
    tf_n6 = tb_n6.text_frame
    tf_n6.margin_left = tf_n6.margin_top = 0
    p_n6 = tf_n6.paragraphs[0]
    apply_style(p_n6, size=10.5, color=COLOR_MUTED)
    p_n6.text = "※ Target 값은 실제 측정값이 아닙니다. Pilot 결과에 따라 수정될 수 있습니다."

    # -------------------------------------------------------------
    # SLIDE 7: 07 / BUSINESS DIRECTION
    # -------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    add_slide_header(s7, "07 / BUSINESS DIRECTION", "첫 제품에서 시작해 기능성 제품으로 확장",
                     "확정된 사업계획이 아니라, 첫 제품의 검증 이후 함께 발전시킬 장기 방향입니다.", "07 / 08")

    stages_3 = [
        ("01", "GLIDE-SPEC 40", "첫 제품 검증\n초기 생산 목표: 3,000개\n\n• 마라토너/행군 타깃의 명확한 효능 입증\n• 극단적 마찰 환경에서 코어 팬덤 구축"),
        ("02", "Sports / Outdoor Functional", "기능성 라인업 확장\n러닝 · 등산 · 사이클 · 군용\n\n• 땀/수분 제어, 쿨링, 피부 보호 제형\n• 구축된 배합 데이터베이스 재활용"),
        ("03", "Data-based Development", "재사용 가능한 개발 구조\n포뮬레이션 플랫폼 자산화\n\n• 원료-공정-QC 데이터베이스 자산화\n• 신제품 개발 주기 및 R&D 비용 획기적 절감")
    ]

    w_s7 = Inches(3.75)
    gap_s7 = Inches(0.24)
    start_s7 = Inches(0.8)
    y_s7 = Inches(2.3)
    h_s7 = Inches(3.7)

    for i, (tag, title, body) in enumerate(stages_3):
        x = start_s7 + i * (w_s7 + gap_s7)
        card = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y_s7, w_s7, h_s7)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_BORDER

        tb = s7.shapes.add_textbox(x + Inches(0.25), y_s7 + Inches(0.22), w_s7 - Inches(0.5), h_s7 - Inches(0.44))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = 0

        p = tf.paragraphs[0]
        apply_style(p, size=12, bold=True, color=COLOR_ACCENT)
        p.text = tag
        p.space_after = Pt(4)

        p_t = tf.add_paragraph()
        apply_style(p_t, size=16, bold=True, color=COLOR_MAIN)
        p_t.text = title
        p_t.space_after = Pt(12)

        p_b = tf.add_paragraph()
        apply_style(p_b, size=12, color=COLOR_BODY)
        p_b.text = body

    # Bottom notice
    tb_n7 = s7.shapes.add_textbox(Inches(0.8), Inches(6.35), Inches(11.7), Inches(0.35))
    tf_n7 = tb_n7.text_frame
    tf_n7.margin_left = tf_n7.margin_top = 0
    p_n7 = tf_n7.paragraphs[0]
    apply_style(p_n7, size=11, color=COLOR_MUTED)
    p_n7.text = "시장 규모나 매출 전망은 아직 넣지 않습니다. 첫 제품의 실제 제조·사용·판매 검증이 먼저입니다."

    # -------------------------------------------------------------
    # SLIDE 8: 08 / COFOUNDER
    # -------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    add_slide_header(s8, "08 / COFOUNDER", "지금부터 필요한 것은 실제 실행",
                     "현재 기술 개발 기준과 데이터 구조를 실제 제품과 사업으로 연결할 파트너가 필요합니다.", "08 / 08")

    w_s8 = Inches(5.7)
    y_s8 = Inches(2.25)
    h_s8 = Inches(3.8)

    # Left: 필요한 실행 역량
    card_l8 = s8.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), y_s8, w_s8, h_s8)
    card_l8.fill.solid()
    card_l8.fill.fore_color.rgb = COLOR_CARD
    card_l8.line.color.rgb = COLOR_BORDER

    tb_l8 = s8.shapes.add_textbox(Inches(1.1), y_s8 + Inches(0.25), w_s8 - Inches(0.6), h_s8 - Inches(0.5))
    tf_l8 = tb_l8.text_frame
    tf_l8.word_wrap = True
    p = tf_l8.paragraphs[0]
    apply_style(p, size=15, bold=True, color=COLOR_MAIN)
    p.text = "필요한 실행 역량"
    p.space_after = Pt(14)

    needs8 = [
        "제조사 / 원료사 네트워크 (소통 및 견적 협상)",
        "Pilot 제조 및 샘플 현장 관리",
        "사용자 필드 테스트 / 시장 검증 (러닝 크루 / 군 장병)",
        "초기 판매 및 유통 채널 (크라우드펀딩 / D2C)",
        "사업 운영 및 전체 마일스톤 일정 관리"
    ]
    for item in needs8:
        pi = tf_l8.add_paragraph()
        apply_style(pi, size=12.5, color=COLOR_BODY)
        pi.text = f"•  {item}"
        pi.space_after = Pt(8)

    # Right: 함께 결정할 것
    card_r8 = s8.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), y_s8, w_s8, h_s8)
    card_r8.fill.solid()
    card_r8.fill.fore_color.rgb = COLOR_CARD
    card_r8.line.color.rgb = COLOR_BORDER

    tb_r8 = s8.shapes.add_textbox(Inches(7.1), y_s8 + Inches(0.25), w_s8 - Inches(0.6), h_s8 - Inches(0.5))
    tf_r8 = tb_r8.text_frame
    tf_r8.word_wrap = True
    p = tf_r8.paragraphs[0]
    apply_style(p, size=15, bold=True, color=COLOR_MAIN)
    p.text = "함께 결정할 것"
    p.space_after = Pt(14)

    decisions8 = [
        "역할 분담 (R&D/시스템 vs 사업개발/운영)",
        "제조 방식 (자체 연구소 조제 vs 외주 OEM)",
        "초기 자금과 예산 배분 계획",
        "초기 시장 진입 방식 (러너 커뮤니티 vs 군수/아웃도어)",
        "첫 생산(3,000개 목표) 및 판매 계획 확정"
    ]
    for item in decisions8:
        pi = tf_r8.add_paragraph()
        apply_style(pi, size=12.5, color=COLOR_BODY)
        pi.text = f"•  {item}"
        pi.space_after = Pt(8)

    # Closing Callout
    card_close = s8.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.25), Inches(11.733), Inches(0.6))
    card_close.fill.solid()
    card_close.fill.fore_color.rgb = COLOR_ACCENT_BG
    card_close.line.color.rgb = COLOR_BORDER

    tb_close = s8.shapes.add_textbox(Inches(1.0), Inches(6.35), Inches(11.333), Inches(0.4))
    tf_close = tb_close.text_frame
    tf_close.margin_left = tf_close.margin_top = 0
    p_close = tf_close.paragraphs[0]
    apply_style(p_close, size=13.5, bold=True, color=COLOR_MAIN, align=PP_ALIGN.CENTER)
    p_close.text = "“이제 이 프로젝트를 실제 제품과 사업으로 만들어가는 단계입니다.”"

    prs.save(OUTPUT_PPTX)
    print(f"Perfect presentation generated at: {OUTPUT_PPTX}")

if __name__ == "__main__":
    create_deck()
