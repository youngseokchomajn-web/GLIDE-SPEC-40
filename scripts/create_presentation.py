#!/usr/bin/env python3
"""
GLIDE-SPEC 40 Presentation Generator
Creates a clean, minimal, restrained 8-slide pitch deck for prospective co-founders.
Adheres strictly to minimal design rules: white/light backgrounds, charcoal/black typography,
ample whitespace, no AI cliches, objective tone.
"""

import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_presentation(output_path: str):
    prs = Presentation()
    # 16:9 Widescreen standard dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]

    # Color Palette: Minimal, restrained, modern typography-first
    COLOR_BG = RGBColor(255, 255, 255)         # Pure White
    COLOR_CARD_BG = RGBColor(250, 250, 250)    # Ultra light neutral grey
    COLOR_BORDER = RGBColor(229, 231, 235)     # Subtle divider line
    COLOR_TEXT_MAIN = RGBColor(17, 24, 39)     # Deep Charcoal (#111827)
    COLOR_TEXT_BODY = RGBColor(55, 65, 81)     # Slate Body (#374151)
    COLOR_TEXT_MUTED = RGBColor(107, 114, 128) # Neutral Grey (#6B7280)
    COLOR_TEXT_LIGHT = RGBColor(156, 163, 175) # Light Meta (#9CA3AF)
    COLOR_ACCENT = RGBColor(30, 58, 138)       # Restrained Deep Slate Navy (#1E3A8A)
    COLOR_ACCENT_BG = RGBColor(241, 245, 249)  # Light slate tint (#F1F5F9)
    COLOR_TAG_BG = RGBColor(243, 244, 246)     # Pill tag background

    FONT_NAME = "Pretendard"
    FALLBACK_FONT = "Apple SD Gothic Neo"

    def apply_font(p, name=FONT_NAME, size=14, bold=False, color=COLOR_TEXT_BODY, align=PP_ALIGN.LEFT):
        p.font.name = name
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
        p.alignment = align

    def add_header(slide, slide_num, title, subtitle=None):
        # Top meta tag
        meta_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.55), Inches(11.333), Inches(0.4))
        tf_meta = meta_box.text_frame
        tf_meta.word_wrap = True
        tf_meta.margin_left = tf_meta.margin_top = tf_meta.margin_right = tf_meta.margin_bottom = 0
        p_meta = tf_meta.paragraphs[0]
        p_meta.text = f"GLIDE-SPEC 40   |   0{slide_num} / 08"
        apply_font(p_meta, size=10, bold=True, color=COLOR_TEXT_LIGHT)

        # Title
        title_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.9), Inches(11.333), Inches(0.7))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title
        apply_font(p_title, size=24, bold=True, color=COLOR_TEXT_MAIN)

        # Subtitle if present
        if subtitle:
            sub_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.55), Inches(11.333), Inches(0.4))
            tf_sub = sub_box.text_frame
            tf_sub.word_wrap = True
            tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
            p_sub = tf_sub.paragraphs[0]
            p_sub.text = subtitle
            apply_font(p_sub, size=13, bold=False, color=COLOR_TEXT_MUTED)

        # Subtle divider
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(2.0), Inches(11.333), Inches(0.01))
        line.fill.solid()
        line.fill.fore_color.rgb = COLOR_BORDER
        line.line.color.rgb = COLOR_BORDER

        # Footer
        footer_box = slide.shapes.add_textbox(Inches(1.0), Inches(6.9), Inches(11.333), Inches(0.3))
        tf_footer = footer_box.text_frame
        tf_footer.margin_left = tf_footer.margin_top = tf_footer.margin_right = tf_footer.margin_bottom = 0
        p_footer = tf_footer.paragraphs[0]
        p_footer.text = "GLIDE-SPEC 40  —  Co-founder Discussion Baseline"
        apply_font(p_footer, size=9, bold=False, color=COLOR_TEXT_LIGHT)

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide (Minimal, Restrained, High Whitespace)
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_slide_layout)
    
    # Category tag
    tag_box = s1.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(10.0), Inches(0.4))
    tf_tag = tag_box.text_frame
    tf_tag.margin_left = tf_tag.margin_top = tf_tag.margin_right = tf_tag.margin_bottom = 0
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = "PROJECT INTRODUCTION"
    apply_font(p_tag, size=11, bold=True, color=COLOR_TEXT_MUTED)

    # Main Project Title
    t_box = s1.shapes.add_textbox(Inches(1.2), Inches(2.4), Inches(10.0), Inches(1.1))
    tf_t = t_box.text_frame
    tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
    p_t = tf_t.paragraphs[0]
    p_t.text = "GLIDE-SPEC 40"
    apply_font(p_t, size=46, bold=True, color=COLOR_TEXT_MAIN)

    # Subtitle
    sub_box = s1.shapes.add_textbox(Inches(1.2), Inches(3.6), Inches(10.0), Inches(0.5))
    tf_sub = sub_box.text_frame
    tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = "Technical Anti-Chafing Stick"
    apply_font(p_sub, size=20, bold=False, color=COLOR_TEXT_BODY)

    # Thin Divider Line
    line1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(4.3), Inches(6.0), Inches(0.015))
    line1.fill.solid()
    line1.fill.fore_color.rgb = COLOR_BORDER
    line1.line.color.rgb = COLOR_BORDER

    # Product Spec Info
    info_box = s1.shapes.add_textbox(Inches(1.2), Inches(4.6), Inches(10.0), Inches(1.2))
    tf_info = info_box.text_frame
    tf_info.margin_left = tf_info.margin_top = tf_info.margin_right = tf_info.margin_bottom = 0
    
    p1 = tf_info.paragraphs[0]
    p1.text = "20g Powder-in-Balm Stick"
    apply_font(p1, size=15, bold=True, color=COLOR_TEXT_BODY)
    p1.space_after = Pt(8)

    p2 = tf_info.add_paragraph()
    p2.text = "Marathon  ·  Trail Running  ·  Long-distance  ·  Military Ruck March"
    apply_font(p2, size=14, bold=False, color=COLOR_TEXT_MUTED)

    # Slide 1 Footer meta
    s1_footer = s1.shapes.add_textbox(Inches(1.2), Inches(6.6), Inches(10.0), Inches(0.4))
    tf_s1_f = s1_footer.text_frame
    tf_s1_f.margin_left = tf_s1_f.margin_top = tf_s1_f.margin_right = tf_s1_f.margin_bottom = 0
    p_s1_f = tf_s1_f.paragraphs[0]
    p_s1_f.text = "Development Baseline: Rev.7.3  |  Engineering Status: Phase C Pilot Preparation"
    apply_font(p_s1_f, size=10, bold=False, color=COLOR_TEXT_LIGHT)


    # -------------------------------------------------------------
    # SLIDE 2: 무엇을 만들고 있는가
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_slide_layout)
    add_header(s2, 2, "무엇을 만들고 있는가", "장시간 고강도 활동 시 발생하는 피부 마찰 문제를 해결하는 기능성 스틱")

    lead_box = s2.shapes.add_textbox(Inches(1.0), Inches(2.25), Inches(11.333), Inches(0.7))
    tf_lead = lead_box.text_frame
    tf_lead.word_wrap = True
    tf_lead.margin_left = tf_lead.margin_top = tf_lead.margin_right = tf_lead.margin_bottom = 0
    p_lead = tf_lead.paragraphs[0]
    p_lead.text = "GLIDE-SPEC 40은 장시간 달리거나 걷는 상황에서 발생하는 피부 마찰을 줄이기 위한 스틱형 제품입니다.\n의류나 피부 간의 반복 마찰로 인한 쓸림(Chafing)을 방지하여 기동력과 집중력을 보존합니다."
    apply_font(p_lead, size=14, bold=False, color=COLOR_TEXT_BODY)

    # Column 1: 주요 사용 부위
    col1_bg = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(3.2), Inches(5.4), Inches(3.2))
    col1_bg.fill.solid()
    col1_bg.fill.fore_color.rgb = COLOR_CARD_BG
    col1_bg.line.color.rgb = COLOR_BORDER
    col1_bg.line.width = Pt(1)

    c1_box = s2.shapes.add_textbox(Inches(1.3), Inches(3.45), Inches(4.8), Inches(2.7))
    tf_c1 = c1_box.text_frame
    tf_c1.word_wrap = True
    tf_c1.margin_left = tf_c1.margin_top = tf_c1.margin_right = tf_c1.margin_bottom = 0
    
    p_c1_title = tf_c1.paragraphs[0]
    p_c1_title.text = "주요 사용 부위"
    apply_font(p_c1_title, size=16, bold=True, color=COLOR_TEXT_MAIN)
    p_c1_title.space_after = Pt(14)

    areas = [
        ("허벅지", "다리 간 지속 마찰로 가장 빈번한 찰과상 발생 부위"),
        ("사타구니", "땀과 열이 집중되어 심한 쓸림과 통증이 발생하는 부위"),
        ("겨드랑이", "팔의 반복 스윙 및 배낭 스트랩 압박 마찰 부위"),
        ("발 & 뒤꿈치", "신발 내 전단 하중으로 물집과 피부 탈락이 발생하는 부위")
    ]
    for name, desc in areas:
        p_item = tf_c1.add_paragraph()
        p_item.text = f"•  {name}  —  {desc}"
        apply_font(p_item, size=13, bold=False, color=COLOR_TEXT_BODY)
        p_item.space_after = Pt(6)

    # Column 2: 핵심 개발 방향
    col2_bg = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.9), Inches(3.2), Inches(5.4), Inches(3.2))
    col2_bg.fill.solid()
    col2_bg.fill.fore_color.rgb = COLOR_CARD_BG
    col2_bg.line.color.rgb = COLOR_BORDER
    col2_bg.line.width = Pt(1)

    c2_box = s2.shapes.add_textbox(Inches(7.2), Inches(3.45), Inches(4.8), Inches(2.7))
    tf_c2 = c2_box.text_frame
    tf_c2.word_wrap = True
    tf_c2.margin_left = tf_c2.margin_top = tf_c2.margin_right = tf_c2.margin_bottom = 0
    
    p_c2_title = tf_c2.paragraphs[0]
    p_c2_title.text = "핵심 개발 방향"
    apply_font(p_c2_title, size=16, bold=True, color=COLOR_TEXT_MAIN)
    p_c2_title.space_after = Pt(14)

    directions = [
        ("마찰 감소", "전단 미끄럼성을 확보하여 동적 마찰 계수 최소화"),
        ("지속성", "땀과 수분에 쉽게 씻겨나가지 않는 견고한 필름 유지"),
        ("사용감", "도포 직후 번들거림이나 끈적임 없는 보송한 마무리"),
        ("전이 최소화", "의류, 양말, 군복으로 유분과 성분이 묻어나지 않도록 제어")
    ]
    for name, desc in directions:
        p_item = tf_c2.add_paragraph()
        p_item.text = f"•  {name}  —  {desc}"
        apply_font(p_item, size=13, bold=False, color=COLOR_TEXT_BODY)
        p_item.space_after = Pt(6)


    # -------------------------------------------------------------
    # SLIDE 3: 제품의 개발 방향
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_slide_layout)
    add_header(s3, 3, "제품의 개발 방향", "Rev.7.3 개발 기준선 — 4대 기능 시스템 구조")

    lead3 = s3.shapes.add_textbox(Inches(1.0), Inches(2.25), Inches(11.333), Inches(0.4))
    tf_lead3 = lead3.text_frame
    tf_lead3.margin_left = tf_lead3.margin_top = tf_lead3.margin_right = tf_lead3.margin_bottom = 0
    p_l3 = tf_lead3.paragraphs[0]
    p_l3.text = "단순 유분막 형성에 의존하지 않고, 4개의 상호 보완적인 화학 시스템으로 분업화된 구조를 설계했습니다."
    apply_font(p_l3, size=14, bold=False, color=COLOR_TEXT_BODY)

    card_w = Inches(2.65)
    card_gap = Inches(0.24)
    start_x = Inches(1.0)
    top_y = Inches(2.85)
    card_h = Inches(3.3)

    systems = [
        ("Powder System", "마찰 저감 & 표면감", 
         "구상 실리카 및 기능성 파우더.\n마이크로 볼베어링 역할을 수행하여 마찰 저항을 줄이고, 땀과 유분을 흡착해 보송한 건식 표면감을 부여합니다."),
        ("Wax System", "스틱 구조 & 사용성", 
         "합성 왁스 및 식물성 왁스 블렌드.\n상온(25°C)에서 단단한 스틱 형상을 지지하고, 피부 도포 시 일정한 양이 균일하게 발리도록 경도를 제어합니다."),
        ("Silicone System", "Spread & Slip", 
         "경량 실리콘 오일 복합체.\n도포 초기 뭉침 없이 피부 전체에 얇고 매끄럽게 펴 발리는 슬립성과 산뜻한 피부 감촉을 형성합니다."),
        ("Resin System", "지속성 & 필름 형성", 
         "MQ Silicone Resin 고형분.\n도포 후 피부 표면에 땀과 수분에 견디는 보호막을 형성하여 장거리 활동 내내 윤활층이 유지되도록 합니다.")
    ]

    for i, (sys_name, role, detail) in enumerate(systems):
        x = start_x + i * (card_w + card_gap)
        bg = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, top_y, card_w, card_h)
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_CARD_BG
        bg.line.color.rgb = COLOR_BORDER
        bg.line.width = Pt(1)

        tb = s3.shapes.add_textbox(x + Inches(0.2), top_y + Inches(0.25), card_w - Inches(0.4), card_h - Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_name = tf.paragraphs[0]
        p_name.text = sys_name
        apply_font(p_name, size=15, bold=True, color=COLOR_TEXT_MAIN)
        p_name.space_after = Pt(4)

        p_role = tf.add_paragraph()
        p_role.text = f"→ {role}"
        apply_font(p_role, size=12, bold=True, color=COLOR_ACCENT)
        p_role.space_after = Pt(14)

        p_det = tf.add_paragraph()
        p_det.text = detail
        apply_font(p_det, size=12, bold=False, color=COLOR_TEXT_BODY)

    note_box = s3.shapes.add_textbox(Inches(1.0), Inches(6.35), Inches(11.333), Inches(0.3))
    tf_note = note_box.text_frame
    tf_note.margin_left = tf_note.margin_top = tf_note.margin_right = tf_note.margin_bottom = 0
    p_n = tf_note.paragraphs[0]
    p_n.text = "* 세부 원료 전체 목록과 개별 배합비는 본 자료에서 생략하며, 각 시스템별 기능적 역할과 물성 제어에 집중합니다."
    apply_font(p_n, size=11, bold=False, color=COLOR_TEXT_MUTED)


    # -------------------------------------------------------------
    # SLIDE 4: 우리가 다른 방식으로 개발하는 부분
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_slide_layout)
    add_header(s4, 4, "제품보다 개발 과정에 더 큰 차이가 있다", "기록과 통계적 검증을 기반으로 한 제품 개발 프로세스")

    flow_y = Inches(2.5)
    steps = ["원료 규격", "배합 설계", "Pilot 제조", "QC 실측", "실제 데이터", "통계적 검증", "최적화"]
    box_w = Inches(1.35)
    box_h = Inches(0.7)
    gap = Inches(0.3)
    start_flow_x = Inches(1.0)

    for i, step in enumerate(steps):
        bx = start_flow_x + i * (box_w + gap)
        box = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, bx, flow_y, box_w, box_h)
        box.fill.solid()
        if i == len(steps) - 1:
            box.fill.fore_color.rgb = COLOR_ACCENT_BG
            box.line.color.rgb = COLOR_ACCENT
            text_color = COLOR_ACCENT
        else:
            box.fill.fore_color.rgb = COLOR_CARD_BG
            box.line.color.rgb = COLOR_BORDER
            text_color = COLOR_TEXT_MAIN
        box.line.width = Pt(1)

        p = box.text_frame.paragraphs[0]
        p.text = step
        apply_font(p, size=12, bold=True, color=text_color, align=PP_ALIGN.CENTER)

        if i < len(steps) - 1:
            arrow_box = s4.shapes.add_textbox(bx + box_w, flow_y + Inches(0.15), gap, Inches(0.4))
            tf_arr = arrow_box.text_frame
            tf_arr.margin_left = tf_arr.margin_top = tf_arr.margin_right = tf_arr.margin_bottom = 0
            p_arr = tf_arr.paragraphs[0]
            p_arr.text = "→"
            apply_font(p_arr, size=16, bold=True, color=COLOR_TEXT_LIGHT, align=PP_ALIGN.CENTER)

    comp_top = Inches(3.65)
    comp_w = Inches(5.45)
    comp_h = Inches(2.7)

    # Conventional
    conv_bg = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), comp_top, comp_w, comp_h)
    conv_bg.fill.solid()
    conv_bg.fill.fore_color.rgb = COLOR_CARD_BG
    conv_bg.line.color.rgb = COLOR_BORDER
    conv_bg.line.width = Pt(1)

    tb_conv = s4.shapes.add_textbox(Inches(1.3), comp_top + Inches(0.3), comp_w - Inches(0.6), comp_h - Inches(0.6))
    tf_conv = tb_conv.text_frame
    tf_conv.word_wrap = True
    tf_conv.margin_left = tf_conv.margin_top = tf_conv.margin_right = tf_conv.margin_bottom = 0
    p_ct = tf_conv.paragraphs[0]
    p_ct.text = "기존 화장품/스틱 개발 방식"
    apply_font(p_ct, size=15, bold=True, color=COLOR_TEXT_MAIN)
    p_ct.space_after = Pt(10)

    p_cb1 = tf_conv.add_paragraph()
    p_cb1.text = "•  경험과 주관적 촉감 중심의 단순 시행착오(Trial & Error) 반복"
    apply_font(p_cb1, size=13, color=COLOR_TEXT_BODY)
    p_cb1.space_after = Pt(6)

    p_cb2 = tf_conv.add_paragraph()
    p_cb2.text = "•  제조 조건(온도, 교반)과 품질 측정 데이터가 체계적으로 남지 않음"
    apply_font(p_cb2, size=13, color=COLOR_TEXT_BODY)
    p_cb2.space_after = Pt(6)

    p_cb3 = tf_conv.add_paragraph()
    p_cb3.text = "•  배치 간 품질 편차가 발생했을 때 원인 규명과 수정이 어려움"
    apply_font(p_cb3, size=13, color=COLOR_TEXT_BODY)

    # Our Way
    our_bg = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.85), comp_top, comp_w, comp_h)
    our_bg.fill.solid()
    our_bg.fill.fore_color.rgb = COLOR_ACCENT_BG
    our_bg.line.color.rgb = COLOR_ACCENT
    our_bg.line.width = Pt(1)

    tb_our = s4.shapes.add_textbox(Inches(7.15), comp_top + Inches(0.3), comp_w - Inches(0.6), comp_h - Inches(0.6))
    tf_our = tb_our.text_frame
    tf_our.word_wrap = True
    tf_our.margin_left = tf_our.margin_top = tf_our.margin_right = tf_our.margin_bottom = 0
    p_ot = tf_our.paragraphs[0]
    p_ot.text = "GLIDE-SPEC 40의 엔지니어링 개발 방식"
    apply_font(p_ot, size=15, bold=True, color=COLOR_ACCENT)
    p_ot.space_after = Pt(10)

    p_ob1 = tf_our.add_paragraph()
    p_ob1.text = "•  투입 원료의 Lot 실측치와 제조 공정 변수(온도/속도) 전수 기록"
    apply_font(p_ob1, size=13, color=COLOR_TEXT_BODY)
    p_ob1.space_after = Pt(6)

    p_ob2 = tf_our.add_paragraph()
    p_ob2.text = "•  경도, 전이도, 용융 거동 등 정량적 QC 지표의 통계적 검증"
    apply_font(p_ob2, size=13, color=COLOR_TEXT_BODY)
    p_ob2.space_after = Pt(6)

    p_ob3 = tf_our.add_paragraph()
    p_ob3.text = "•  축적된 실측 데이터를 기반으로 다음 포뮬러를 체계적으로 최적화"
    apply_font(p_ob3, size=13, color=COLOR_TEXT_BODY)


    # -------------------------------------------------------------
    # SLIDE 5: 현재 상태
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_slide_layout)
    add_header(s5, 5, "현재는 실제 Pilot을 준비하는 단계", "엔지니어링 프레임워크 구축 완료 및 물리적 제조 진입 직전")

    col_y = Inches(2.3)
    col_h = Inches(3.4)

    # Column 1: 완료된 사항
    bg_done = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), col_y, Inches(5.5), col_h)
    bg_done.fill.solid()
    bg_done.fill.fore_color.rgb = COLOR_CARD_BG
    bg_done.line.color.rgb = COLOR_BORDER
    bg_done.line.width = Pt(1)

    tb_done = s5.shapes.add_textbox(Inches(1.25), col_y + Inches(0.25), Inches(5.0), col_h - Inches(0.5))
    tf_done = tb_done.text_frame
    tf_done.word_wrap = True
    tf_done.margin_left = tf_done.margin_top = tf_done.margin_right = tf_done.margin_bottom = 0

    p_dt = tf_done.paragraphs[0]
    p_dt.text = "완료된 엔지니어링 자산 (Completed)"
    apply_font(p_dt, size=15, bold=True, color=COLOR_TEXT_MAIN)
    p_dt.space_after = Pt(10)

    done_items = [
        "제품 개발 기준 확정 (Rev.7.3 Baseline)",
        "원료 규격 및 배합 데이터 구조 확립",
        "QC 측정 지표 및 품질 데이터 구조 표준화",
        "실험계획법(DOE) 및 통계 분석 Framework",
        "포뮬러 최적화 및 적격성 평가(Qualification) 엔진",
        "코드 및 배합 형상 관리(Revision Tracking)",
        "기초 시스템 테스트 31/31 Unit Tests Pass"
    ]
    for item in done_items:
        pi = tf_done.add_paragraph()
        pi.text = f"•  {item}"
        apply_font(pi, size=12, color=COLOR_TEXT_BODY)
        pi.space_after = Pt(4)

    # Column 2: 현재 단계 및 다음 단계
    bg_now = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), col_y, Inches(5.5), col_h)
    bg_now.fill.solid()
    bg_now.fill.fore_color.rgb = COLOR_CARD_BG
    bg_now.line.color.rgb = COLOR_BORDER
    bg_now.line.width = Pt(1)

    tb_now = s5.shapes.add_textbox(Inches(7.05), col_y + Inches(0.25), Inches(5.0), col_h - Inches(0.5))
    tf_now = tb_now.text_frame
    tf_now.word_wrap = True
    tf_now.margin_left = tf_now.margin_top = tf_now.margin_right = tf_now.margin_bottom = 0

    p_nt = tf_now.paragraphs[0]
    p_nt.text = "현재 위치 (Current Focus)"
    apply_font(p_nt, size=15, bold=True, color=COLOR_ACCENT)
    p_nt.space_after = Pt(8)

    p_curr = tf_now.add_paragraph()
    p_curr.text = "Phase C — Physical Pilot Execution"
    apply_font(p_curr, size=14, bold=True, color=COLOR_TEXT_MAIN)
    p_curr.space_after = Pt(4)

    p_curr_desc = tf_now.add_paragraph()
    p_curr_desc.text = "이론적 계산 및 데이터베이스 구축을 마치고, 실제 실험실 환경에서 원료를 계량·제조하는 실행 단계로 진입합니다."
    apply_font(p_curr_desc, size=12, color=COLOR_TEXT_MUTED)
    p_curr_desc.space_after = Pt(16)

    p_next_t = tf_now.add_paragraph()
    p_next_t.text = "직전 다음 단계 (Immediate Next)"
    apply_font(p_next_t, size=14, bold=True, color=COLOR_TEXT_MAIN)
    p_next_t.space_after = Pt(8)

    p_next_flow = tf_now.add_paragraph()
    p_next_flow.text = "실제 원료 입고  →  Pilot 제조 (18 Runs)  →  QC 실측  →  통계적 Qualification"
    apply_font(p_next_flow, size=12, bold=True, color=COLOR_TEXT_BODY)

    # Critical Warning Box at bottom
    warn_bg = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(5.9), Inches(11.3), Inches(0.85))
    warn_bg.fill.solid()
    warn_bg.fill.fore_color.rgb = COLOR_ACCENT_BG
    warn_bg.line.color.rgb = COLOR_BORDER
    warn_bg.line.width = Pt(1)

    tb_w = s5.shapes.add_textbox(Inches(1.2), Inches(6.0), Inches(10.9), Inches(0.65))
    tf_w = tb_w.text_frame
    tf_w.word_wrap = True
    tf_w.margin_left = tf_w.margin_top = tf_w.margin_right = tf_w.margin_bottom = 0
    p_wt = tf_w.paragraphs[0]
    p_wt.text = "중요 원칙: Production Model은 아직 Qualified(공식 자격 승격) 상태가 아닙니다."
    apply_font(p_wt, size=13, bold=True, color=COLOR_TEXT_MAIN)
    p_wt.space_after = Pt(2)

    p_wd = tf_w.add_paragraph()
    p_wd.text = "공개 벤치마크 데이터로 알고리즘 안정성은 검증했으나, 실제 제품 자격은 향후 파일럿 실측 데이터가 확보되어야 최종 판정됩니다."
    apply_font(p_wd, size=11, color=COLOR_TEXT_MUTED)


    # -------------------------------------------------------------
    # SLIDE 6: 앞으로 해야 할 일
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_slide_layout)
    add_header(s6, 6, "이제부터는 실제 실행이 중요하다", "실제 파일럿 제조부터 초기 양산 검토까지의 8단계 로드맵")

    row_h = Inches(1.85)
    col_w = Inches(2.65)
    gap_x = Inches(0.24)
    start_x = Inches(1.0)
    
    steps_roadmap = [
        ("01", "원료 공급사 확보", "핵심 원료(Resin, Wax, Silicone 등)의 공급선 컨택, CoA 및 샘플 입고"),
        ("02", "제조/Pilot 파트너 확보", "소량 파일럿 조제 및 시험 생산이 가능한 OEM/연구 파트너 섭외"),
        ("03", "18-run Pilot 실행", "무작위 제조 순서와 중심점 4회 반복이 설계된 18개 시험 배치 제조"),
        ("04", "QC 및 제조 데이터 확보", "경도(gf), 도포 전이량(g), 온도 편차 등 정량적 물성 실측치 계측"),
        ("05", "실제 사용자 테스트", "마라톤 풀코스, 트레일러너, 군장 행군 대상자 필드 쓸림 방어 테스트"),
        ("06", "결과 분석 및 최적화", "실측치 회귀 분석 및 한계 영역 파악을 통한 Formula Refinement"),
        ("07", "Confirmation Run", "최종 확정 포뮬러에 대한 재현성 검증 및 생산 규격 승격"),
        ("08", "초기 생산 검토", "초기 목표 3,000개 생산 준비 (Target COGS ~₩2,950 / unit)")
    ]

    for idx, (num, stitle, sdesc) in enumerate(steps_roadmap):
        r = idx // 4
        c = idx % 4
        x = start_x + c * (col_w + gap_x)
        y = Inches(2.35) + r * (row_h + Inches(0.35))

        bg_step = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, col_w, row_h)
        bg_step.fill.solid()
        bg_step.fill.fore_color.rgb = COLOR_CARD_BG
        bg_step.line.color.rgb = COLOR_BORDER
        bg_step.line.width = Pt(1)

        tb_step = s6.shapes.add_textbox(x + Inches(0.2), y + Inches(0.18), col_w - Inches(0.4), row_h - Inches(0.3))
        tf_step = tb_step.text_frame
        tf_step.word_wrap = True
        tf_step.margin_left = tf_step.margin_top = tf_step.margin_right = tf_step.margin_bottom = 0

        p_num = tf_step.paragraphs[0]
        p_num.text = num
        apply_font(p_num, size=11, bold=True, color=COLOR_ACCENT)
        p_num.space_after = Pt(2)

        p_st = tf_step.add_paragraph()
        p_st.text = stitle
        apply_font(p_st, size=14, bold=True, color=COLOR_TEXT_MAIN)
        p_st.space_after = Pt(6)

        p_sd = tf_step.add_paragraph()
        p_sd.text = sdesc
        apply_font(p_sd, size=11, color=COLOR_TEXT_MUTED)


    # -------------------------------------------------------------
    # SLIDE 7: 사업으로서의 방향
    # -------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_slide_layout)
    add_header(s7, 7, "첫 제품은 GLIDE-SPEC 40", "단일 제품에서 시작하여 기능성 제품 라인업과 개발 플랫폼으로의 확장")

    block_w = Inches(3.6)
    block_gap = Inches(0.26)
    b_y = Inches(2.4)
    b_h = Inches(3.0)

    stages = [
        ("Step 1", "GLIDE-SPEC 40", 
         "초경량 테크니컬 안티체이핑 스틱\n\n• 마라토너/장거리 행군 타깃 검증\n• 극단적 마찰 환경에서의 명확한 효능 입증\n• 확실한 효능을 통한 초기 코어 팬덤 구축\n• 초기 생산 목표: 3,000개"),
        ("Step 2", "Sports / Outdoor Functional", 
         "스포츠·아웃도어 기능성 라인업 확장\n\n• 땀/수분 제어 기능성 제품군\n• 쿨링 및 피부 보호 제형 확장\n• 러닝, 등산, 사이클, 군용 맞춤 스펙\n• 데이터 기반 배합 시스템 재활용"),
        ("Step 3", "데이터 기반 제품 개발 시스템", 
         "시행착오를 줄이는 포뮬레이션 플랫폼\n\n• 원료-공정-QC 데이터베이스 자산화\n• 신제품 개발 주기 및 R&D 비용 획기적 단축\n• 반복 가능하고 재현성 높은 B2B/브랜드 기반\n• 제조 데이터 기반의 진입 장벽 구축")
    ]

    for i, (tag, title, body) in enumerate(stages):
        bx = start_x + i * (block_w + block_gap)
        bg = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, bx, b_y, block_w, b_h)
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_CARD_BG
        bg.line.color.rgb = COLOR_BORDER
        bg.line.width = Pt(1)

        tb = s7.shapes.add_textbox(bx + Inches(0.25), b_y + Inches(0.25), block_w - Inches(0.5), b_h - Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_tag = tf.paragraphs[0]
        p_tag.text = tag
        apply_font(p_tag, size=11, bold=True, color=COLOR_ACCENT)
        p_tag.space_after = Pt(4)

        p_t = tf.add_paragraph()
        p_t.text = title
        apply_font(p_t, size=16, bold=True, color=COLOR_TEXT_MAIN)
        p_t.space_after = Pt(12)

        p_b = tf.add_paragraph()
        p_b.text = body
        apply_font(p_b, size=12, color=COLOR_TEXT_BODY)

    pnote_box = s7.shapes.add_textbox(Inches(1.0), Inches(5.75), Inches(11.333), Inches(0.9))
    tf_pn = pnote_box.text_frame
    tf_pn.word_wrap = True
    tf_pn.margin_left = tf_pn.margin_top = tf_pn.margin_right = tf_pn.margin_bottom = 0

    p_pn1 = tf_pn.paragraphs[0]
    p_pn1.text = "장기적으로는 Anti-Chafing 제품 하나에 머무르지 않고, 스포츠·아웃도어용 기능성 제품으로 확장할 수 있는 기반을 만드는 것이 목표입니다."
    apply_font(p_pn1, size=13, bold=True, color=COLOR_TEXT_MAIN)
    p_pn1.space_after = Pt(4)

    p_pn2 = tf_pn.add_paragraph()
    p_pn2.text = "* 본 내용은 과장된 사업 계획서가 아니며, 첫 번째 제품을 성공적으로 론칭한 이후 함께 논의하며 발전시켜 나갈 장기적 방향성입니다."
    apply_font(p_pn2, size=11, color=COLOR_TEXT_MUTED)


    # -------------------------------------------------------------
    # SLIDE 8: 공동창업자와 논의하고 싶은 것
    # -------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_slide_layout)
    add_header(s8, 8, "지금 필요한 것은 실제 실행이다", "엔지니어링 설계를 바탕으로 함께 만들어갈 실행 과제")

    lead8 = s8.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.333), Inches(0.5))
    tf_l8 = lead8.text_frame
    tf_l8.word_wrap = True
    tf_l8.margin_left = tf_l8.margin_top = tf_l8.margin_right = tf_l8.margin_bottom = 0
    p_l8 = tf_l8.paragraphs[0]
    p_l8.text = "현재 기술 개발 기준과 데이터 분석 구조는 상당 부분 준비되어 있습니다. 이제 책상 위의 계획을 실제 제품과 사업으로 완성할 파트너가 필요합니다."
    apply_font(p_l8, size=14, bold=False, color=COLOR_TEXT_BODY)

    col8_w = Inches(5.45)
    col8_h = Inches(3.0)

    # Left: 앞으로 필요한 실제 실행 역량
    b_left = s8.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(2.85), col8_w, col8_h)
    b_left.fill.solid()
    b_left.fill.fore_color.rgb = COLOR_CARD_BG
    b_left.line.color.rgb = COLOR_BORDER
    b_left.line.width = Pt(1)

    tb_l = s8.shapes.add_textbox(Inches(1.3), Inches(3.1), col8_w - Inches(0.6), col8_h - Inches(0.5))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True
    tf_l.margin_left = tf_l.margin_top = tf_l.margin_right = tf_l.margin_bottom = 0

    p_lt = tf_l.paragraphs[0]
    p_lt.text = "앞으로 필요한 실제 실행"
    apply_font(p_lt, size=16, bold=True, color=COLOR_TEXT_MAIN)
    p_lt.space_after = Pt(12)

    needs = [
        "제조사 및 원료사 네트워크 (OEM/ODM 소통 및 협상)",
        "실제 Pilot 제조 현장 실행 및 샘플 관리",
        "사용자 필드 테스트 진행 (러닝 크루 / 군 장병)",
        "타깃 시장 검증 및 고객 피드백 정량화",
        "초기 판매 및 유통 채널 확보 (크라우드펀딩 / D2C)",
        "사업 운영 전반 및 일정 관리"
    ]
    for need in needs:
        pn = tf_l.add_paragraph()
        pn.text = f"•  {need}"
        apply_font(pn, size=12, color=COLOR_TEXT_BODY)
        pn.space_after = Pt(4)

    # Right: 함께 결정할 사항
    b_right = s8.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.85), Inches(2.85), col8_w, col8_h)
    b_right.fill.solid()
    b_right.fill.fore_color.rgb = COLOR_CARD_BG
    b_right.line.color.rgb = COLOR_BORDER
    b_right.line.width = Pt(1)

    tb_r = s8.shapes.add_textbox(Inches(7.15), Inches(3.1), col8_w - Inches(0.6), col8_h - Inches(0.5))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True
    tf_r.margin_left = tf_r.margin_top = tf_r.margin_right = tf_r.margin_bottom = 0

    p_rt = tf_r.paragraphs[0]
    p_rt.text = "함께 결정할 사항"
    apply_font(p_rt, size=16, bold=True, color=COLOR_TEXT_MAIN)
    p_rt.space_after = Pt(12)

    decisions = [
        "역할 분담 (R&D/시스템 vs 사업개발/운영)",
        "제조 방식 (자체 파일럿 연구소 조제 vs 외주 OEM 진행)",
        "초기 자금 조달 및 예산 배분 계획",
        "초기 시장 진입 전략 (러너 커뮤니티 vs 군수/아웃도어)",
        "첫 생산(3,000개) 및 론칭 타임라인 확정",
        "지분 및 파트너십 구조"
    ]
    for d in decisions:
        pd = tf_r.add_paragraph()
        pd.text = f"•  {d}"
        apply_font(pd, size=12, color=COLOR_TEXT_BODY)
        pd.space_after = Pt(4)

    # Bottom Closing Statement
    closing_bg = s8.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(6.05), Inches(11.3), Inches(0.7))
    closing_bg.fill.solid()
    closing_bg.fill.fore_color.rgb = COLOR_ACCENT_BG
    closing_bg.line.color.rgb = COLOR_BORDER
    closing_bg.line.width = Pt(1)

    tb_c = s8.shapes.add_textbox(Inches(1.2), Inches(6.2), Inches(10.9), Inches(0.4))
    tf_c = tb_c.text_frame
    tf_c.margin_left = tf_c.margin_top = tf_c.margin_right = tf_c.margin_bottom = 0
    p_ctext = tf_c.paragraphs[0]
    p_ctext.text = "“이제 이 프로젝트를 실제 제품과 사업으로 만들어가는 단계입니다.”"
    apply_font(p_ctext, size=14, bold=True, color=COLOR_TEXT_MAIN, align=PP_ALIGN.CENTER)

    # Save
    prs.save(output_path)
    print(f"Presentation successfully created at: {output_path}")

if __name__ == "__main__":
    output = "docs/presentation/GLIDE_SPEC_40_CoFounder_Intro.pptx"
    if len(sys.argv) > 1:
        output = sys.argv[1]
    build_presentation(output)
